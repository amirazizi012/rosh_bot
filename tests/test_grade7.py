"""آزمون پایه‌ی هفتم: همه‌ی امکانات تکلیف، بدون نشان و امتیاز."""
import os, sys, tempfile
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.dirname(__file__))
from app import rules as R, reports, panel_ui as UI, points, scheduler, db
from test_flow import make_app, upd, go, ADMIN

S6, S7, S9 = 3000, 3500, 3900
H, D = 3600, 86400


def mk():
    app, bale = make_app()
    go(app, upd(ADMIN, "/start"),
       upd(ADMIN, data="a:sadd"), upd(ADMIN, data="a:sadd:6"), upd(ADMIN, "علی ششمی"),
       upd(ADMIN, data="a:sadd:7"), upd(ADMIN, "رضا هفتمی"),
       upd(ADMIN, data="a:sadd:9"), upd(ADMIN, "نیما نهمی"))
    for chat, name in ((S6, "علی ششمی"), (S7, "رضا هفتمی"), (S9, "نیما نهمی")):
        code = app.conn.execute("SELECT code FROM students WHERE full_name=?", (name,)).fetchone()["code"]
        go(app, upd(chat, "/start"), upd(chat, code))
    app.conn.execute("UPDATE students SET created_ts=?", (app._t - 90 * D,)); app.conn.commit()
    return app, bale


def sid(app, name):
    return app.conn.execute("SELECT id FROM students WHERE full_name=?", (name,)).fetchone()["id"]


def add_asg(app, title, kind="required", dl=None, target_all=1, target_grade=0, points=0):
    dl = dl or app._t + 5 * H
    app.conn.execute("INSERT INTO assignments(title,kind,deadline_ts,grace_min,late_hours,target_all,created_ts,points,target_grade) VALUES(?,?,?,20,24,?,?,?,?)",
                     (title, kind, dl, target_all, app._t - 6 * D, points, target_grade))
    app.conn.commit()
    return app.conn.execute("SELECT * FROM assignments ORDER BY id DESC").fetchone()


def submit_db(app, a, s, at):
    cur = app.conn.execute("INSERT INTO submissions(assignment_id,student_id,state,created_ts) VALUES(?,?,'draft',?)", (a["id"], s, at))
    app.conn.commit()
    return R.finalize(app.conn, a, s, cur.lastrowid, at)


def test_add_grade7_and_menus():
    app, bale = mk()
    go(app, upd(ADMIN, data="a:sadd"))
    labels = [b["text"] for row in bale.out[-1][2] for b in row]
    assert "پایه‌ی هفتم" in labels and "پایه‌ی ششم" in labels and "پایه‌ی نهم" in labels
    s7 = app.conn.execute("SELECT * FROM students WHERE full_name='رضا هفتمی'").fetchone()
    assert s7["grade"] == 7 and app.student_by_chat(S7)
    # منوی هفتم: بدون نشان؛ منوی ششم: با نشان
    go(app, upd(S7, "/start"))
    cb7 = [b["text"] for row in bale.out[-1][2] for b in row]
    assert "📈 پیشرفت من" in cb7 and not any("نشان" in t for t in cb7) and "📚 تکالیف باز" in cb7
    go(app, upd(S6, "/start"))
    assert any("🏅 نشان‌ها و پیشرفت" in b["text"] for row in bale.out[-1][2] for b in row)
    # فهرست مدیر پایه را نشان می‌دهد
    go(app, upd(ADMIN, data="a:slist:0"))
    assert any("رضا هفتمی (هفتم)" in b["text"] for row in bale.out[-1][2] for b in row)


def test_assignment_targeting_by_grade():
    app, bale = mk()
    s6, s7, s9 = sid(app, "علی ششمی"), sid(app, "رضا هفتمی"), sid(app, "نیما نهمی")
    ids = lambda a: sorted(s["id"] for s in R.target_students(app.conn, a))
    both, only6, only7 = add_asg(app, "هر دو"), add_asg(app, "فقط ششم", target_grade=6), add_asg(app, "فقط هفتم", target_grade=7)
    assert ids(both) == [s6, s7] and ids(only6) == [s6] and ids(only7) == [s7]
    assert [a["title"] for a in R.targets_for_student(app.conn, s7)] == ["هر دو", "فقط هفتم"]
    assert R.targets_for_student(app.conn, s9) == []
    # تکلیف گروهی: هر عضو پایه‌ی ششم یا هفتم دارد
    app.conn.execute("INSERT INTO groups(name) VALUES('تقویتی')"); app.conn.execute("INSERT INTO student_groups VALUES(?,1)", (s7,)); app.conn.commit()
    g = add_asg(app, "گروهی", target_all=0)
    app.conn.execute("INSERT INTO assignment_groups VALUES(?,1)", (g["id"],)); app.conn.commit()
    assert ids(g) == [s7]
    # مسیر مدیر: ساخت تکلیف فقط برای هفتم
    go(app, upd(ADMIN, data="a:new"), upd(ADMIN, "هندسه"), upd(ADMIN, "حل کنید"), upd(ADMIN, data="a:nb"), upd(ADMIN, data="a:nk:required"))
    labels = [b["text"] for row in bale.out[-1][2] for b in row]
    assert any("هر دو پایه (ششم و هفتم)" in l for l in labels) and "فقط ششم" in labels and "فقط هفتم" in labels
    go(app, upd(ADMIN, data="a:nt:g7"), upd(ADMIN, "فردا ۲۱:۳۰"))
    assert "همه‌ی پایه‌ی هفتم" in bale.texts(ADMIN)[-1]
    go(app, upd(ADMIN, data="a:npub:1"))
    a = app.conn.execute("SELECT * FROM assignments ORDER BY id DESC").fetchone()
    assert (a["target_all"], a["target_grade"]) == (1, 7)
    assert any("هندسه" in t for t in bale.texts(S7)) and not any("هندسه" in t for t in bale.texts(S6))
    # دانش‌آموز پایه‌ی دیگر با شناسه‌ی دستی هم نمی‌تواند آن را باز کند
    go(app, upd(S6, data=f"s:a:{a['id']}"))
    assert "پیدا نشد" in bale.texts(S6)[-1]
    # گزینه‌ی «هر دو پایه» (سازگار با قبل)
    go(app, upd(ADMIN, data="a:new"), upd(ADMIN, "مشترک"), upd(ADMIN, "x"), upd(ADMIN, data="a:nb"), upd(ADMIN, data="a:nk:required"),
       upd(ADMIN, data="a:nt:all"), upd(ADMIN, "فردا ۲۱:۳۰"), upd(ADMIN, data="a:npub:0"))
    a2 = app.conn.execute("SELECT * FROM assignments ORDER BY id DESC").fetchone()
    assert (a2["target_all"], a2["target_grade"]) == (1, 0) and ids(a2) == [s6, s7]
    go(app, upd(ADMIN, data=f"a:as:{a2['id']}"))
    assert "هر دو پایه (ششم و هفتم)" in bale.texts(ADMIN)[-1]


def test_grade7_has_all_homework_features():
    app, bale = mk()
    s7 = sid(app, "رضا هفتمی")
    a = add_asg(app, "ریاضی", dl=app._t + 3 * H, target_grade=7)
    # ارسال با دو فایل، ویرایش تا مهلت
    go(app, upd(S7, data=f"s:send:{a['id']}"), upd(S7, photo=True), upd(S7, photo=True), upd(S7, data="s:fin"))
    f1 = R.current_final(app.conn, a["id"], s7)
    assert f1["on_time"] == 1
    go(app, upd(S7, data=f"s:ed:{a['id']}"))
    fid = app.conn.execute("SELECT id FROM submission_files WHERE submission_id=(SELECT id FROM submissions WHERE state='draft')").fetchone()["id"]
    go(app, upd(S7, data=f"s:rm:{fid}"), upd(S7, data="s:fin"))
    assert R.current_final(app.conn, a["id"], s7)["id"] != f1["id"]
    # مدیر: ناقص ← دانش‌آموز پیام می‌گیرد و دوباره می‌فرستد
    go(app, upd(ADMIN, data=f"a:rvs:{a['id']}:{s7}:inc"), upd(ADMIN, "ناقص بود"))
    assert R.student_status(app.conn, a, s7, app._t)["code"] == "incomplete" and any("«ناقص»" in t for t in bale.texts(S7))
    go(app, upd(S7, data=f"s:send:{a['id']}"), upd(S7, photo=True), upd(S7, data="s:fin"))
    assert R.student_status(app.conn, a, s7, app._t)["code"] == "on_time"
    # تاخیر و بسته‌شدن
    b = add_asg(app, "فیزیک", dl=app._t - 2 * H, target_grade=7)
    go(app, upd(S7, data=f"s:send:{b['id']}"), upd(S7, photo=True), upd(S7, data="s:fin"))
    assert "۲ ساعت" in bale.texts(S7)[-3] + bale.texts(S7)[-2] + bale.texts(S7)[-1] and R.current_final(app.conn, b["id"], s7)["on_time"] == 0
    c = add_asg(app, "ادبیات", dl=app._t - 30 * H, target_grade=7)
    go(app, upd(S7, data=f"s:send:{c['id']}"))
    assert "بسته" in bale.texts(S7)[-1]
    # عذر موجه، تمدید فردی، بازخورد، یادآوری، پیام همگانی
    go(app, upd(ADMIN, data=f"a:aexcs:{c['id']}:{s7}"), upd(ADMIN, "بیماری"))
    assert R.is_excused(app.conn, c["id"], s7)
    d = add_asg(app, "شیمی", dl=app._t + 2 * H, target_grade=7)
    scheduler.reminders(app)
    assert any("یادآوری" in t and "شیمی" in t for t in bale.texts(S7))
    go(app, upd(ADMIN, data=f"a:fbs:{s7}:{d['id']}"), upd(ADMIN, "آفرین"))
    assert any("آفرین" in t for t in bale.texts(S7))
    go(app, upd(ADMIN, data="a:bc"), upd(ADMIN, data="a:bct:g7"), upd(ADMIN, "فردا کلاس جبرانی"), upd(ADMIN, data="a:bcsend"))
    assert any("کلاس جبرانی" in t for t in bale.texts(S7)) and not any("کلاس جبرانی" in t for t in bale.texts(S6))
    # ارسال‌های من
    go(app, upd(S7, data="s:done"))
    assert "ریاضی" in bale.texts(S7)[-1]
    # ZIP پاسخ‌ها
    go(app, upd(ADMIN, data=f"a:azip:{a['id']}"))
    assert any(f.endswith(".zip") for _, f in bale.files)


def test_no_badges_or_points_for_grade7():
    app, bale = mk()
    s6, s7 = sid(app, "علی ششمی"), sid(app, "رضا هفتمی")
    # ۶ الزامی + ۵ اختیاری + یک امتیازی با امتیاز، همه به‌موقع برای هر دو پایه
    for i in range(6):
        a = add_asg(app, f"الزامی {i}", dl=app._t - (20 - i) * D)
        for s in (s6, s7):
            submit_db(app, a, s, a["deadline_ts"] - H)
    for i in range(5):
        a = add_asg(app, f"اختیاری {i}", kind="optional", dl=app._t - (10 - i) * D)
        for s in (s6, s7):
            submit_db(app, a, s, a["deadline_ts"] - H)
    bonus = add_asg(app, "ویژه", kind="bonus", dl=app._t + 4 * H, points=15)
    # ششم نشان/امتیاز در انتظار می‌گیرد؛ هفتم هیچ
    assert app.award(s6) and app.award(s7) == []
    assert R.give_badge(app.conn, s7, "manual", "1", "x", app._t) is False
    assert R.add_points(app.conn, s7, 5, "x", "manual", None, None, app._t, "approved") is False
    # مدیر: ارسال امتیازی هفتمی ← امتیازی ثبت نمی‌شود؛ ششمی ← در انتظار تأیید
    go(app, upd(S7, data=f"s:a:{bonus['id']}"))
    assert "بعد از تأیید مدیر" not in bale.texts(S7)[-1]                           # در متن تکلیف امتیاز نشان داده نمی‌شود
    go(app, upd(S7, data=f"s:send:{bonus['id']}"), upd(S7, photo=True), upd(S7, data="s:fin"))
    assert not any("تکلیف امتیازی است" in t for t in bale.texts(S7))
    go(app, upd(S6, data=f"s:a:{bonus['id']}"))
    assert "⭐ ۱۵ امتیاز" in bale.texts(S6)[-1]
    go(app, upd(S6, data=f"s:send:{bonus['id']}"), upd(S6, photo=True), upd(S6, data="s:fin"))
    assert any("تکلیف امتیازی است" in t for t in bale.texts(S6))
    assert app.conn.execute("SELECT COUNT(*) c FROM badges WHERE student_id=?", (s7,)).fetchone()["c"] == 0
    assert app.conn.execute("SELECT COUNT(*) c FROM points_ledger WHERE student_id=?", (s7,)).fetchone()["c"] == 0
    assert app.conn.execute("SELECT COUNT(*) c FROM points_ledger WHERE student_id=? AND kind='bonus'", (s6,)).fetchone()["c"] == 1
    # صف تأیید فقط موردهای ششم را دارد
    names = [r["full_name"] for r in app.conn.execute("SELECT s.full_name FROM badges b JOIN students s ON s.id=b.student_id WHERE b.status='pending'")]
    assert names and set(names) == {"علی ششمی"}
    # بررسی وضعیت برای هفتم کار می‌کند ولی چیزی ثبت نمی‌کند
    go(app, upd(ADMIN, data=f"a:rvs:{bonus['id']}:{s7}:ok"))
    assert app.conn.execute("SELECT COUNT(*) c FROM points_ledger WHERE student_id=?", (s7,)).fetchone()["c"] == 0
    # دکمه‌ها و صفحه‌های مدیر
    go(app, upd(ADMIN, data=f"a:st:{s7}"))
    btns = [b["callback_data"] for row in bale.out[-1][2] for b in row]
    assert f"a:rvl:{s7}" in btns and f"a:sbd:{s7}" not in btns and f"a:spts:{s7}" not in btns and f"a:stgrp:{s7}" in btns
    go(app, upd(ADMIN, data=f"a:st:{s6}"))
    btns = [b["callback_data"] for row in bale.out[-1][2] for b in row]
    assert f"a:sbd:{s6}" in btns and f"a:spts:{s6}" in btns
    go(app, upd(ADMIN, data=f"a:sbd:{s7}"))
    assert "فعال نیست" in bale.texts(ADMIN)[-1]
    go(app, upd(ADMIN, data=f"a:spts:{s7}"))
    assert "فعال نیست" in bale.texts(ADMIN)[-1]
    go(app, upd(ADMIN, data="a:badge"))
    pick = [b["text"] for row in bale.out[-1][2] for b in row]
    assert "علی ششمی" in pick and "رضا هفتمی" not in pick
    # پیشرفت دانش‌آموز هفتم: آمار بدون نشان و امتیاز
    go(app, upd(S7, data="s:prog"))
    t = bale.texts(S7)[-1]
    assert "پیشرفت تو" in t and "به‌موقع" in t and "نشان" not in t and "امتیاز تو" not in t
    go(app, upd(S6, data="s:prog"))
    assert "امتیاز تو" in bale.texts(S6)[-1]


def test_reports_and_panel_for_grade7():
    app, bale = mk()
    s6, s7 = sid(app, "علی ششمی"), sid(app, "رضا هفتمی")
    a = add_asg(app, "الف", dl=app._t - 2 * D)
    submit_db(app, a, s6, a["deadline_ts"] - H)
    submit_db(app, a, s7, a["deadline_ts"] + 3 * H)
    txt = reports.class_lines(app)
    assert "پایه‌ی هفتم (بدون نشان و امتیاز)" in txt and "رضا هفتمی" in txt and "نیما نهمی" not in txt
    line7 = next(l for l in txt.splitlines() if l.startswith("رضا هفتمی"))
    assert "⭐" not in line7 and "🏅" not in line7
    rep7 = reports.student_report(app, s7)
    assert "پایه‌ی هفتم" in rep7 and "امتیاز تأییدشده" not in rep7 and "نشان‌ها" not in rep7 and "⭐" not in rep7
    rep6 = reports.student_report(app, s6)
    assert "امتیاز تأییدشده" in rep6 and "نشان‌ها" in rep6
    # اکسل
    from openpyxl import load_workbook
    p = reports.build_xlsx(app, app.data / "exports" / "t.xlsx")
    ws = load_workbook(p).worksheets[0]
    head = [c.value for c in ws[1]]
    rows = {r[0].value: [c.value for c in r] for r in ws.iter_rows(min_row=2)}
    assert "پایه" in head and rows["رضا هفتمی"][1] == "هفتم" and rows["علی ششمی"][1] == "ششم"
    assert rows["رضا هفتمی"][head.index("امتیاز تأییدشده")] == "—" and rows["علی ششمی"][head.index("امتیاز تأییدشده")] != "—"
    # پنل
    dash = UI.dashboard(app.conn, app.cfg, app._t)
    assert "رضا هفتمی" in dash and "(هفتم)" in dash and "علی ششمی" in dash
    pg7 = UI.student_page(app.conn, app.cfg, app._t, s7)
    assert "پایه‌ی هفتم" in pg7 and "مسیر نشان بعدی" not in pg7 and "دفتر امتیاز" not in pg7 and "امتیاز تأییدشده" not in pg7
    pg6 = UI.student_page(app.conn, app.cfg, app._t, s6)
    assert "مسیر نشان بعدی" in pg6 and "دفتر امتیاز" in pg6
    assert "هفتم" in UI.students_page(app.conn, app.cfg, app._t)
    assert "رضا هفتمی" not in UI.points_page(app.conn, app.cfg, app._t) and "رضا هفتمی" not in UI.badges_page(app.conn, app.cfg, app._t)
    b = add_asg(app, "فقط هفتم", target_grade=7)
    assert "همه‌ی هفتم" in UI.assignments_page(app.conn, app.cfg, app._t) and "همه‌ی هفتم" in UI.assignment_page(app.conn, app.cfg, app._t, b["id"])


def test_old_assignments_stay_for_grade6_only_after_migration():
    import sqlite3
    p = tempfile.mktemp(suffix=".db")
    c = sqlite3.connect(p)
    c.executescript("CREATE TABLE students(id INTEGER PRIMARY KEY AUTOINCREMENT, full_name TEXT NOT NULL, code TEXT UNIQUE NOT NULL, chat_id INTEGER UNIQUE, active INTEGER NOT NULL DEFAULT 1, created_ts INTEGER NOT NULL, joined_ts INTEGER);"
                    "INSERT INTO students(full_name,code,created_ts) VALUES('ششمی قدیمی','1',1);"
                    "CREATE TABLE assignments(id INTEGER PRIMARY KEY AUTOINCREMENT, title TEXT NOT NULL, description TEXT NOT NULL DEFAULT '', kind TEXT NOT NULL DEFAULT 'required', deadline_ts INTEGER NOT NULL, grace_min INTEGER NOT NULL, late_hours INTEGER NOT NULL, target_all INTEGER NOT NULL DEFAULT 1, status TEXT NOT NULL DEFAULT 'active', created_ts INTEGER NOT NULL);"
                    "INSERT INTO assignments(title,deadline_ts,grace_min,late_hours,created_ts) VALUES('قدیمی',9999999999,20,24,1);")
    c.commit(); c.close()
    conn = db.connect(p)
    conn.execute("INSERT INTO students(full_name,code,created_ts,grade) VALUES('هفتمی تازه','2',1,7)"); conn.commit()
    assert conn.execute("SELECT target_grade FROM assignments").fetchone()["target_grade"] == 6
    assert len(R.targets_for_student(conn, 1)) == 1 and R.targets_for_student(conn, 2) == []       # تکلیف قبلی فقط برای ششم


if __name__ == "__main__":
    for k, v in list(globals().items()):
        if k.startswith("test_"):
            v(); print("✓", k)
