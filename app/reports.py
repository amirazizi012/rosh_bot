"""گزارش‌ها: اکسل کامل، متن گزارش، فایل ZIP پاسخ‌ها."""
from __future__ import annotations

import re
import zipfile
from pathlib import Path

from . import rules as R
from . import jalali as J
from .fa import fa, duration

ZIP_PART_LIMIT = 40 * 1024 * 1024  # هر بسته حداکثر ~۴۰ مگابایت


def safe_name(s: str) -> str:
    return re.sub(r'[\\/:*?"<>|\n\r]+', "_", s).strip() or "بی‌نام"


def cell_text(st: dict) -> str:
    code = st["code"]
    if code == "on_time":
        return "✅ به‌موقع"
    if code == "late":
        return f"⏰ {duration(st['over'])} تاخیر"
    if code == "excused":
        return "🙏 عذر موجه"
    if code == "waiting":
        return "… در انتظار"
    return "❌ ارسال‌نشده"


def group_names(conn, sid) -> str:
    rows = conn.execute("SELECT g.name FROM groups g JOIN student_groups sg ON sg.group_id=g.id "
                        "WHERE sg.student_id=? ORDER BY g.name", (sid,)).fetchall()
    return "، ".join(r["name"] for r in rows)


def class_lines(app) -> str:
    now, n = app.now(), app.cfg.streak_size
    studs = app.conn.execute("SELECT * FROM students WHERE active=1 AND grade=6 ORDER BY full_name").fetchall()
    if not studs:
        return "هنوز دانش‌آموز پایه‌ی ششمی ثبت نشده است."
    lines = [f"📊 خلاصه‌ی تکالیف الزامی پایه‌ی ششم ({J.fmt_short(now)})",
             "✅ به‌موقع | ⏰ با تاخیر | ❌ ارسال‌نشده | 🙏 عذر | 🔥 شمارنده | 🏅 نشان | ⭐ امتیاز (در انتظار)", ""]
    for s in studs:
        t = R.student_stats(app.conn, s["id"], now, n)
        lines.append(f"{s['full_name']}: ✅{fa(t['on_time'])} ⏰{fa(t['late'])} ❌{fa(t['missed'])} "
                     f"🙏{fa(t['excused'])} | 🔥{fa(t['current'])}/{fa(n)} | 🏅{fa(t['badges'])} | ⭐{fa(t['points'])}"
                     + (f" ({fa(t['points_pending'])})" if t["points_pending"] else ""))
    return "\n".join(lines)


def student_report(app, sid) -> str:
    conn, now, n = app.conn, app.now(), app.cfg.streak_size
    s = app.student(sid)
    if s["grade"] == 9:
        return student_book_report(conn, s)
    t = R.student_stats(conn, sid, now, n)
    lines = [f"👤 {s['full_name']}", f"گروه‌ها: {group_names(conn, sid) or '—'}", "",
             f"تکالیف الزامی: ✅ {fa(t['on_time'])} به‌موقع | ⏰ {fa(t['late'])} با تاخیر | "
             f"❌ {fa(t['missed'])} ارسال‌نشده | 🙏 {fa(t['excused'])} عذر موجه",
             f"شمارنده‌ی فعلی: {fa(t['current'])} از {fa(n)} | اختیاری/امتیازی انجام‌شده: {fa(t['extra_done'])}",
             f"⭐ امتیاز تأییدشده: {fa(t['points'])} | در انتظار تأیید: {fa(t['points_pending'])}", ""]
    lines.append("جزئیات تکالیف:")
    for a in reversed(R.targets_for_student(conn, sid)):
        st = R.student_status(conn, a, sid, now)
        extra = ""
        if st["sub"]:
            extra = f" — ارسال: {J.fmt_short(st['sub']['submitted_ts'])}"
        lines.append(f"{R.KIND_ICON[a['kind']]} {a['title']}: {cell_text(st)}{extra}")
    bd = conn.execute("SELECT title FROM badges WHERE student_id=? ORDER BY awarded_ts", (sid,)).fetchall()
    lines += ["", "نشان‌ها:"] + ([f"• {b['title']}" for b in bd] or ["—"])
    fb = conn.execute("SELECT COUNT(*) c FROM feedback WHERE student_id=?", (sid,)).fetchone()["c"]
    lines.append(f"\nبازخوردهای ثبت‌شده: {fa(fb)}")
    return "\n".join(lines)


def build_xlsx(app, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    conn, now, n = app.conn, app.now(), app.cfg.streak_size
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="2E7D4F")
    head_font = Font(name="Tahoma", bold=True, color="FFFFFF")
    body_font = Font(name="Tahoma")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)

    def style(ws, header_row=1):
        ws.sheet_view.rightToLeft = True
        for row in ws.iter_rows():
            for c in row:
                c.alignment = center
                c.font = head_font if c.row == header_row else body_font
                if c.row == header_row:
                    c.fill = head_fill
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 3), 38)

    studs = conn.execute("SELECT * FROM students WHERE active=1 AND grade=6 ORDER BY full_name").fetchall()
    ws = wb.active
    ws.title = "خلاصه"
    ws.append(["نام", "گروه‌ها", "الزامی: به‌موقع", "الزامی: با تاخیر", "الزامی: ارسال‌نشده",
               "عذر موجه", "شمارنده‌ی فعلی", "نشان‌های پشتکار", "تعداد کل نشان‌ها", "اختیاری/امتیازی",
               "امتیاز تأییدشده", "امتیاز در انتظار تأیید"])
    for s in studs:
        t = R.student_stats(conn, s["id"], now, n)
        ws.append([s["full_name"], group_names(conn, s["id"]), fa(t["on_time"]), fa(t["late"]), fa(t["missed"]),
                   fa(t["excused"]), f"{fa(t['current'])} از {fa(n)}", fa(t["blocks"]), fa(t["badges"]),
                   fa(t["extra_done"]), fa(t["points"]), fa(t["points_pending"])])
    style(ws)

    asgs = conn.execute("SELECT * FROM assignments WHERE status='active' ORDER BY deadline_ts, id").fetchall()
    ws2 = wb.create_sheet("ماتریس تکالیف")
    ws2.append(["نام"] + [f"{a['title']}\n{R.KIND_LABEL[a['kind']]} | {J.fmt_short(a['deadline_ts'])}" for a in asgs])
    for s in studs:
        row = [s["full_name"]]
        mine = {a["id"] for a in R.targets_for_student(conn, s["id"])}
        for a in asgs:
            row.append(cell_text(R.student_status(conn, a, s["id"], now)) if a["id"] in mine else "—")
        ws2.append(row)
    style(ws2)
    ws2.row_dimensions[1].height = 48

    ws3 = wb.create_sheet("بازخوردها")
    ws3.append(["دانش‌آموز", "تکلیف", "زمان", "متن بازخورد"])
    for r in conn.execute("SELECT f.*, s.full_name, a.title FROM feedback f JOIN students s ON s.id=f.student_id "
                          "LEFT JOIN assignments a ON a.id=f.assignment_id ORDER BY f.ts"):
        ws3.append([r["full_name"], r["title"] or "—", J.fmt_short(r["ts"]), r["text"]])
    style(ws3)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _final_files(conn, aid=None, sid=None):
    sql = ("SELECT sf.path, sf.name, s.assignment_id, s.student_id, st.full_name, a.title, sf.id fid "
           "FROM submission_files sf JOIN submissions s ON s.id=sf.submission_id "
           "JOIN students st ON st.id=s.student_id JOIN assignments a ON a.id=s.assignment_id "
           "WHERE s.state='final'")
    args = []
    if aid is not None:
        sql += " AND s.assignment_id=?"
        args.append(aid)
    if sid is not None:
        sql += " AND s.student_id=?"
        args.append(sid)
    return conn.execute(sql + " ORDER BY st.full_name, s.assignment_id, sf.id", args).fetchall()


def build_zips(app, out_dir: Path, aid=None, sid=None, label="answers"):
    """فایل‌های نهایی را در یک یا چند ZIP (هر کدام زیر حدود ۴۰MB) می‌ریزد."""
    rows = _final_files(app.conn, aid, sid)
    out_dir.mkdir(parents=True, exist_ok=True)
    parts, cur, cur_size, idx = [], None, 0, 0
    counters = {}
    for r in rows:
        p = Path(r["path"])
        if not p.exists():
            continue
        size = p.stat().st_size
        if cur is None or cur_size + size > ZIP_PART_LIMIT:
            if cur:
                cur.close()
            idx += 1
            zp = out_dir / f"{safe_name(label)}_{idx}.zip"
            cur = zipfile.ZipFile(zp, "w", zipfile.ZIP_STORED)
            parts.append(zp)
            cur_size = 0
        key = (r["student_id"], r["assignment_id"])
        counters[key] = counters.get(key, 0) + 1
        folder = safe_name(r["full_name"]) if sid is None else safe_name(r["title"])
        prefix = safe_name(r["title"]) if sid is None and aid is None else ""
        arc = f"{folder}/{prefix + '_' if prefix else ''}{counters[key]}{p.suffix}"
        if aid is not None and sid is None:
            arc = f"{safe_name(r['full_name'])}_{counters[key]}{p.suffix}"
        cur.write(p, arc)
        cur_size += size
    if cur:
        cur.close()
    return parts


# ====================== باشگاه کتابخوانی (پایه‌ی نهم) ======================
def week_start(ts: int) -> int:
    """شروع هفته (شنبه) به‌صورت ظهر همان روز"""
    d = J.to_dt(ts)
    return J.noon(ts) - ((d.weekday() - 5) % 7) * 86400


def _norm_book(s: str) -> str:
    return "".join(str(s).split()).replace("ي", "ی").replace("ك", "ک").lower()


def book_stats(conn, since=None):
    """آمار کلی کتابخوانی: هر نفر، هر کتاب، هر هفته. (فقط گزارش‌های فعال دانش‌آموزان فعال پایه‌ی نهم)"""
    q = ("SELECT b.*, s.full_name FROM book_reports b JOIN students s ON s.id=b.student_id "
         "WHERE b.status='active' AND s.grade=9 AND s.active=1")
    args = []
    if since:
        q += " AND b.day2_ts>=?"
        args.append(since)
    rows = conn.execute(q + " ORDER BY b.day2_ts, b.id", args).fetchall()
    per = {s["id"]: {"student": s, "reports": 0, "pages": 0, "books": set(), "weeks": set(), "last": 0}
           for s in conn.execute("SELECT * FROM students WHERE active=1 AND grade=9 ORDER BY full_name")}
    books, weekly = {}, {}
    for r in rows:
        p = per.get(r["student_id"])
        if p is None:
            continue
        key = _norm_book(r["book"])
        p["reports"] += 1
        p["pages"] += r["pages"]
        p["books"].add(key)
        w = week_start(r["day2_ts"])
        p["weeks"].add(w)
        p["last"] = max(p["last"], r["day2_ts"])
        b = books.setdefault(key, {"title": r["book"], "pages": 0, "readers": set()})
        b["pages"] += r["pages"]
        b["readers"].add(r["student_id"])
        wk = weekly.setdefault(w, [0, 0])
        wk[0] += r["pages"]
        wk[1] += 1
    students = list(per.values())
    return {"rows": rows, "students": students,
            "books": sorted(books.values(), key=lambda b: -b["pages"]),
            "weekly": sorted(weekly.items()),
            "total_pages": sum(p["pages"] for p in students), "total_reports": sum(p["reports"] for p in students),
            "readers": sum(1 for p in students if p["reports"])}


def book_summary_text(conn, now: int, since=None) -> str:
    st = book_stats(conn, since)
    lines = [f"📖 گزارش کلی باشگاه کتابخوانی ({J.fmt_short(now)})", "",
             f"دانش‌آموزان: {fa(st['readers'])} نفر از {fa(len(st['students']))} نفر گزارش داده‌اند",
             f"گزارش‌ها: {fa(st['total_reports'])} | مجموع صفحه‌ها: {fa(st['total_pages'])} | کتاب‌ها: {fa(len(st['books']))}", "",
             "هر نفر (به ترتیب الفبا):"]
    for p in st["students"]:
        last = J.fmt_short(p["last"]).split(" ")[0] if p["last"] else "—"
        lines.append(f"• {p['student']['full_name']}: {fa(p['reports'])} گزارش | {fa(p['pages'])} صفحه | "
                     f"{fa(len(p['books']))} کتاب | آخرین: {last}")
    if st["books"]:
        lines += ["", "کتاب‌ها:"] + [f"• {b['title']}: {fa(len(b['readers']))} نفر | {fa(b['pages'])} صفحه" for b in st["books"][:15]]
    return "\n".join(lines)


def student_book_report(conn, s) -> str:
    rows = conn.execute("SELECT * FROM book_reports WHERE student_id=? AND status='active' ORDER BY day2_ts DESC, id DESC", (s["id"],)).fetchall()
    tot = sum(r["pages"] for r in rows)
    lines = [f"👤 {s['full_name']} — پایه‌ی نهم", f"گزارش‌ها: {fa(len(rows))} | مجموع صفحه‌ها: {fa(tot)}", ""]
    for r in rows[:40]:
        lines.append(f"📕 {r['book']} — {J.fmt_short(r['day2_ts']).split(' ')[0]} — صفحه‌ی {fa(r['from_page'])} تا {fa(r['to_page'])} ({fa(r['pages'])} صفحه)")
    return "\n".join(lines)


def _xl_style(ws):
    from openpyxl.styles import Alignment, Font, PatternFill
    ws.sheet_view.rightToLeft = True
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, readingOrder=2)
            c.font = Font(name="Tahoma", bold=True, color="FFFFFF") if c.row == 1 else Font(name="Tahoma")
            if c.row == 1:
                c.fill = PatternFill("solid", fgColor="2E7D4F")
    for col in ws.columns:
        w = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(12, w + 3), 40)


def build_books_xlsx(app, path: Path, since=None) -> Path:
    from openpyxl import Workbook
    conn = app.conn
    st = book_stats(conn, since)
    wb = Workbook()
    ws = wb.active
    ws.title = "خلاصه"
    ws.append(["نام", "تعداد گزارش", "مجموع صفحه", "تعداد کتاب", "هفته‌های فعال", "آخرین گزارش"])
    for p in st["students"]:
        ws.append([p["student"]["full_name"], fa(p["reports"]), fa(p["pages"]), fa(len(p["books"])), fa(len(p["weeks"])),
                   J.fmt_short(p["last"]).split(" ")[0] if p["last"] else "—"])
    _xl_style(ws)
    ws2 = wb.create_sheet("گزارش‌ها")
    ws2.append(["نام", "روزهای گزارش", "کتاب", "از صفحه", "تا صفحه", "تعداد صفحه", "زمان ثبت", "ارسال به گروه"])
    for r in reversed(st["rows"]):
        ws2.append([r["full_name"], re.sub(r"\s+", " ", _pair(r)), r["book"], fa(r["from_page"]), fa(r["to_page"]),
                    fa(r["pages"]), J.fmt_short(r["created_ts"]), "بله" if r["posted"] else "خیر"])
    _xl_style(ws2)
    ws3 = wb.create_sheet("کتاب‌ها")
    ws3.append(["کتاب", "تعداد خواننده", "مجموع صفحه"])
    for b in st["books"]:
        ws3.append([b["title"], fa(len(b["readers"])), fa(b["pages"])])
    _xl_style(ws3)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _pair(r) -> str:
    from .books import pair_label
    return pair_label(r["day1_ts"], r["day2_ts"])
